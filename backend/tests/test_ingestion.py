import re
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from backend.app.rag.ingestion import chunk_text, ingest_manual, load_manual, load_manual_bytes

MANUAL_TXT = Path(__file__).parent / "manual_test.txt"


# --- load_manual ---

def test_load_manual_txt_reads_content():
    text = load_manual(MANUAL_TXT)
    assert "Facebook" in text


def test_load_manual_missing_file_raises():
    with pytest.raises(FileNotFoundError):
        load_manual(Path("does_not_exist.txt"))


def test_load_manual_unsupported_extension_raises():
    with pytest.raises(ValueError):
        load_manual(Path(__file__))  # .py ไม่รองรับ


# --- chunk_text ---

def test_chunk_text_empty_returns_empty_list():
    assert chunk_text("") == []


def test_chunk_text_shorter_than_chunk_size_returns_single_chunk():
    assert chunk_text("hello world", chunk_size=500) == ["hello world"]


def test_chunk_text_splits_multiline_text():
    text = "\n".join(f"line {i}" for i in range(100))
    chunks = chunk_text(text, chunk_size=50, overlap=10)
    assert len(chunks) > 1
    for c in chunks:
        assert len(c) <= 50 + 10 + 1  # เผื่อ overlap prefix ต่อท้าย


def test_chunk_text_splits_oversized_single_line():
    # บรรทัดเดียวยาว 1200 ตัวอักษร ไม่มี newline (เช่น PDF ที่ extract ไม่มี newline)
    text = "a" * 1200
    chunks = chunk_text(text, chunk_size=500, overlap=50)
    assert len(chunks) > 1
    assert all(len(c) <= 500 + 50 + 1 for c in chunks)
    assert "".join(chunks).replace(" ", "").count("a") >= 1200


def test_chunk_text_keeps_each_paragraph_as_its_own_chunk():
    # แต่ละหัวข้อ (คั่นด้วยบรรทัดว่าง) ต้องไม่ถูกรวมเข้า chunk เดียวกับหัวข้ออื่น
    # แม้จะรวมกันแล้วยังพอดี chunk_size ก็ตาม (บั๊กเดิม: ยัดหลายหัวข้อใน chunk เดียว)
    text = "topic A: do X\n\ntopic B: do Y\n\ntopic C: do Z"
    chunks = chunk_text(text, chunk_size=500)
    assert chunks == ["topic A: do X", "topic B: do Y", "topic C: do Z"]


def test_chunk_text_real_manual_produces_one_chunk_per_topic():
    text = load_manual(MANUAL_TXT)
    chunks = chunk_text(text)
    # manual_test.txt มี 62 ย่อหน้าคั่นด้วยบรรทัดว่าง (นับรวมหัวข้อใหม่ D-J)
    assert len(chunks) == 62
    assert "Facebook" in chunks[0]
    assert "Alt+F4" in chunks[1]


def test_chunk_text_never_splits_a_word_across_chunks():
    # ย่อหน้าเดียวยาวมาก หลายประโยค ไม่มีบรรทัดว่างคั่น -> ต้อง fallback ไปแบ่งตามประโยค/คำ
    # และห้ามมี chunk ไหนขึ้นต้นหรือจบกลางคำ (บั๊กเดิม: overlap ตัดกลางคำ)
    words = [f"word{i}" for i in range(200)]
    text = ". ".join(" ".join(words[i:i + 8]) for i in range(0, len(words), 8)) + "."
    chunks = chunk_text(text, chunk_size=120, overlap=20)
    assert len(chunks) > 1
    for c in chunks:
        # ทุก token ในทุก chunk ต้องเป็น "wordN" เต็มคำเสมอ ห้ามมีเศษคำที่ถูกตัดครึ่ง
        for token in c.replace(".", " ").split():
            assert re.fullmatch(r"word\d+", token), f"found broken token: {token!r} in chunk {c!r}"


# --- ingest_manual (mock chroma กันยิง DB จริง/โหลดโมเดล embedding ตอนรัน unit test) ---
# หมายเหตุ: ingestion.py ไม่ embed เองแล้ว (local embedding ผูกไว้กับ collection ใน
# chroma_client.py) — ingest_manual แค่ต้องส่ง documents ดิบให้ collection.upsert()

def test_ingest_manual_upserts_with_matching_ids_and_documents():
    with patch("backend.app.rag.ingestion.get_collection") as mock_get_collection, \
         patch("backend.app.rag.ingestion.chunk_text", return_value=["chunk A", "chunk B"]):
        mock_collection = mock_get_collection.return_value
        mock_collection.name = "manuals"

        ingest_manual(MANUAL_TXT)

        mock_collection.upsert.assert_called_once()
        _, kwargs = mock_collection.upsert.call_args
        assert kwargs["documents"] == ["chunk A", "chunk B"]
        assert "embeddings" not in kwargs  # ให้ ChromaDB embed เองด้วย local model
        assert len(kwargs["ids"]) == 2
        assert len(set(kwargs["ids"])) == 2  # id ไม่ชนกัน
        assert all(m["source"] == MANUAL_TXT.name for m in kwargs["metadatas"])


def test_ingest_manual_skips_when_no_chunks():
    with patch("backend.app.rag.ingestion.chunk_text", return_value=[]), \
         patch("backend.app.rag.ingestion.get_collection") as mock_get_collection:
        ingest_manual(MANUAL_TXT)

        mock_get_collection.return_value.upsert.assert_not_called()


# --- edge cases เพิ่มเติม (chunk_text) ---

def test_chunk_text_whitespace_only_returns_empty_list():
    """ข้อความที่มีแต่ whitespace/newline ต้องคืน [] เหมือน empty string"""
    assert chunk_text("   \n\n   \t  \n") == []


def test_chunk_text_single_word_longer_than_chunk_size_does_not_crash():
    """คำเดียวที่ยาวกว่า chunk_size จริงๆ — fallback hard-split ตัวอักษร ไม่ควร crash"""
    long_word = "x" * 1500
    chunks = chunk_text(long_word, chunk_size=500, overlap=0)
    assert len(chunks) > 1
    # ต้องครอบคลุมตัวอักษรครบ
    assert sum(len(c) for c in chunks) >= 1500


def test_chunk_text_overlap_zero_no_prefix_added():
    """overlap=0 ต้องไม่มี prefix ต่อท้ายขึ้นต้น chunk ถัดไป (ไม่มีเศษ overlap)"""
    # ใช้ประโยคสั้นๆ สองย่อหน้า ให้ chunk ที่ 1 จบแล้วไม่ต้อง overlap ไป chunk ที่ 2
    text = "A " * 300   # บรรทัดเดียว ยาวพอให้แตกหลาย chunk
    chunks = chunk_text(text, chunk_size=100, overlap=0)
    assert len(chunks) > 1
    # ทุก chunk ต้องไม่ยาวเกิน chunk_size (ไม่มี overlap เพิ่ม)
    for c in chunks:
        assert len(c) <= 100


def test_chunk_text_new_topics_present_in_real_manual():
    """หัวข้อใหม่ที่เพิ่มเข้า manual_test.txt ต้องถูก chunk แยกและมีเนื้อหาครบ"""
    text = load_manual(MANUAL_TXT)
    chunks = chunk_text(text)
    combined = " ".join(chunks)
    # ตรวจว่าหัวข้อใหม่ทั้ง 5 มีอยู่ใน chunks จริง
    assert "Clear browsing data" in combined   # ล้าง Cache
    assert "Incognito" in combined             # Incognito/Private
    assert "Ctrl+D" in combined                # Bookmark
    assert "Print Preview" in combined         # Print
    assert "Ctrl+0" in combined                # Zoom


def test_ingest_manual_metadata_contains_path_field():
    """metadata ของทุก chunk ต้องมี field 'path' เป็น string (ไม่ใช่ None/missing)"""
    with patch("backend.app.rag.ingestion.get_collection") as mock_get_collection, \
         patch("backend.app.rag.ingestion.chunk_text", return_value=["chunk X"]):
        mock_collection = mock_get_collection.return_value
        mock_collection.name = "manuals"

        ingest_manual(MANUAL_TXT)

        _, kwargs = mock_collection.upsert.call_args
        for meta in kwargs["metadatas"]:
            assert "path" in meta
            assert isinstance(meta["path"], str)


def test_ingest_manual_ids_differ_across_files():
    """ไฟล์ต่างกัน ต้องได้ id ต่างกัน (hash ขึ้นกับ stem ของไฟล์) กัน id ชนข้ามไฟล์"""
    import hashlib
    other_path = MANUAL_TXT.parent / "other_manual.txt"
    id_a = hashlib.md5(f"{MANUAL_TXT.stem}_0".encode()).hexdigest()
    id_b = hashlib.md5(f"{other_path.stem}_0".encode()).hexdigest()
    assert id_a != id_b


# --- edge cases เพิ่มเติม รอบ 2 ---

def test_chunk_text_only_newlines_returns_empty_list():
    """string ที่มีแต่ newlines หลายบรรทัด ต้องคืน [] เหมือน empty"""
    assert chunk_text("\n\n\n\n\n") == []


def test_chunk_text_exactly_chunk_size_is_single_chunk():
    """ข้อความที่ยาวพอดี chunk_size ต้องได้ 1 chunk ไม่ถูกแตกออก"""
    text = "A" * 500
    chunks = chunk_text(text, chunk_size=500, overlap=0)
    assert len(chunks) == 1
    assert chunks[0] == text


def test_chunk_text_overlap_larger_than_piece_does_not_crash():
    """overlap ใหญ่กว่าเนื้อหาของ piece ต้องไม่ crash และคืน list"""
    text = "short sentence. next sentence."
    chunks = chunk_text(text, chunk_size=10, overlap=50)
    assert isinstance(chunks, list)
    assert len(chunks) > 0


def test_chunk_text_mixed_thai_english_preserves_all_content():
    """ข้อความผสม TH/EN ต้องไม่ทำให้เนื้อหาหายหรือบิดเบือน"""
    text = "วิธีใช้ browser: กด Ctrl+T เพื่อเปิด new tab"
    chunks = chunk_text(text, chunk_size=500)
    combined = " ".join(chunks)
    assert "Ctrl+T" in combined
    assert "browser" in combined


def test_chunk_text_group_a_topics_present_in_real_manual():
    """หัวข้อกลุ่ม A (Tab/Reload/DevTools/URL/Navigation) ต้องอยู่ใน chunks จริง"""
    text = load_manual(MANUAL_TXT)
    chunks = chunk_text(text)
    combined = " ".join(chunks)
    assert "Ctrl+T" in combined          # เปิด Tab ใหม่
    assert "Ctrl+Shift+T" in combined    # กู้คืน Tab
    assert "Alt+Left Arrow" in combined  # Browser History
    assert "F5" in combined              # Reload
    assert "F12" in combined             # Developer Tools
    assert "Ctrl+L" in combined          # Copy URL
    assert "Ctrl+Tab" in combined        # สลับ Tab


def test_chunk_text_group_b_topics_present_in_real_manual():
    """หัวข้อกลุ่ม B (Form/Dropdown/Upload/Copy) ต้องอยู่ใน chunks จริง"""
    text = load_manual(MANUAL_TXT)
    chunks = chunk_text(text)
    combined = " ".join(chunks)
    assert "Submit" in combined          # กรอกฟอร์ม
    assert "dropdown" in combined        # Dropdown
    assert "Choose File" in combined     # Upload
    assert "select all" in combined      # คัดลอกข้อความ


def test_chunk_text_group_c_topics_present_in_real_manual():
    """หัวข้อกลุ่ม C (Logout/Password/HTTPS) ต้องอยู่ใน chunks จริง"""
    text = load_manual(MANUAL_TXT)
    chunks = chunk_text(text)
    combined = " ".join(chunks)
    assert "Sign out" in combined        # Logout
    assert "current password" in combined  # เปลี่ยนรหัสผ่าน
    assert "padlock" in combined         # SSL/HTTPS
    assert "https://" in combined        # SSL/HTTPS


def test_ingest_manual_chunk_index_is_sequential():
    """metadata chunk_index ต้องเรียงจาก 0 ถึง N-1 ต่อเนื่องไม่มีช่องว่าง"""
    with patch("backend.app.rag.ingestion.get_collection") as mock_get_collection, \
         patch("backend.app.rag.ingestion.chunk_text",
               return_value=["a", "b", "c", "d", "e"]):
        mock_collection = mock_get_collection.return_value
        mock_collection.name = "manuals"

        ingest_manual(MANUAL_TXT)

        _, kwargs = mock_collection.upsert.call_args
        indices = [m["chunk_index"] for m in kwargs["metadatas"]]
        assert indices == list(range(5))


# --- pdf/xlsx: load_manual_bytes (ต่างจาก load_manual ด้านบนตรงรับ bytes ไม่ต้องมีไฟล์บน
# ดิสก์ — ใช้กับไฟล์ที่ user แนบเข้ามาผ่าน API โดยตรง) ---

def _xlsx_bytes(*sheets: tuple[str, list[list]]) -> bytes:
    """สร้างไฟล์ XLSX จริงในหน่วยความจำ (openpyxl round-trip เต็มรูปแบบ ไม่ mock) —
    sheets: (sheet_name, rows) หลายแผ่นได้ ชื่อ sheet แรกจะทับ default "Sheet" ของ Workbook()"""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_load_manual_bytes_txt_reads_content():
    text = load_manual_bytes("Facebook login page".encode("utf-8"), "note.txt")
    assert text == "Facebook login page"


def test_load_manual_bytes_unsupported_extension_raises():
    with pytest.raises(ValueError):
        load_manual_bytes(b"whatever", "script.py")


def test_load_manual_bytes_pdf_extracts_page_text():
    # mock PdfReader ตรงจุดขอบเขตเดียวกับที่ test_ingest_manual_* ด้านบน mock get_collection
    # (ขอบเขต external library ไม่ใช่ logic ของเราเอง) — สร้าง PDF ไบนารีจริงด้วยมือเปราะบาง
    # เกินไป ไม่มีประโยชน์เพิ่มเทียบกับ mock ตรงนี้
    mock_page = MagicMock()
    mock_page.extract_text.return_value = "Invoice total: 1,250 THB"
    with patch("backend.app.rag.ingestion.PdfReader") as mock_reader_cls:
        mock_reader_cls.return_value.pages = [mock_page]
        text = load_manual_bytes(b"%PDF-1.4 fake bytes", "invoice.pdf")
    assert "Invoice total: 1,250 THB" in text
    mock_reader_cls.assert_called_once()
    # ต้องส่ง BytesIO เข้า PdfReader (ไม่ใช่ path) — คือหัวใจของ "ไม่แตะดิสก์" ของฟังก์ชันนี้
    assert isinstance(mock_reader_cls.call_args[0][0], BytesIO)


def test_load_manual_bytes_xlsx_single_sheet():
    content = _xlsx_bytes(("Data", [["Name", "Price"], ["Widget", 9.99], ["Gadget", 19.99]]))
    text = load_manual_bytes(content, "report.xlsx")
    assert "# Sheet: Data" in text
    assert "Name | Price" in text
    assert "Widget | 9.99" in text
    assert "Gadget | 19.99" in text


def test_load_manual_bytes_xlsx_multiple_sheets():
    content = _xlsx_bytes(
        ("Q1", [["Jan", 100]]),
        ("Q2", [["Apr", 200]]),
    )
    text = load_manual_bytes(content, "quarters.xlsx")
    assert "# Sheet: Q1" in text
    assert "Jan | 100" in text
    assert "# Sheet: Q2" in text
    assert "Apr | 200" in text


def test_load_manual_bytes_xlsx_skips_fully_empty_rows():
    content = _xlsx_bytes(("Data", [["A", "B"], [None, None], ["C", "D"]]))
    text = load_manual_bytes(content, "report.xlsx")
    lines = [line for line in text.split("\n") if line]
    assert lines == ["# Sheet: Data", "A | B", "C | D"]


def test_load_manual_bytes_xlsx_none_cells_become_empty_string():
    content = _xlsx_bytes(("Data", [["A", None, "C"]]))
    text = load_manual_bytes(content, "report.xlsx")
    assert "A |  | C" in text


# Excel extractor: merged-cell forward-fill — openpyxl only stores a value in the
# top-left cell of a merged range; every other cell in that range reads back as None
# even though Excel visually shows the same value there (e.g. an employee name merged
# across several date rows in a timesheet). Sub-rows must not go blank because of this.

def test_load_manual_bytes_xlsx_vertically_merged_cell_forward_fills_to_subrows():
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Date", "Hours"])
    ws.append(["Somchai", "2026-07-01", 8])
    ws.append([None, "2026-07-02", 8])
    ws.append([None, "2026-07-03", 8])
    ws.merge_cells("A2:A4")
    buf = BytesIO()
    wb.save(buf)

    text = load_manual_bytes(buf.getvalue(), "timesheet.xlsx")

    lines = [line for line in text.split("\n") if line]
    assert lines == [
        "# Sheet: Data",
        "Name | Date | Hours",
        "Somchai | 2026-07-01 | 8",
        "Somchai | 2026-07-02 | 8",
        "Somchai | 2026-07-03 | 8",
    ]


def test_load_manual_bytes_xlsx_horizontally_merged_header_forward_fills():
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Department", None, "Total"])
    ws.append(["Engineering", None, 4000])
    ws.merge_cells("A1:B1")
    ws.merge_cells("A2:B2")
    buf = BytesIO()
    wb.save(buf)

    text = load_manual_bytes(buf.getvalue(), "report.xlsx")

    lines = [line for line in text.split("\n") if line]
    assert lines == [
        "# Sheet: Data",
        "Department | Department | Total",
        "Engineering | Engineering | 4000",
    ]


def test_load_manual_bytes_xlsx_merged_cell_does_not_break_empty_row_skip():
    """merge ไม่ควรทำให้แถวที่ว่างจริงๆ (ไม่ได้อยู่ใน merge range ไหนเลย) กลายเป็นแถวที่มี
    ข้อมูลไปเฉยๆ — เช็คว่า _xlsx_bytes_to_text() ยัง skip แถวว่างล้วนๆ ได้ตามปกติ"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Hours"])
    ws.append(["Somchai", 8])
    ws.append([None, None])
    ws.append(["Somsri", 6])
    buf = BytesIO()
    wb.save(buf)

    text = load_manual_bytes(buf.getvalue(), "timesheet.xlsx")

    lines = [line for line in text.split("\n") if line]
    assert lines == ["# Sheet: Data", "Name | Hours", "Somchai | 8", "Somsri | 6"]


# --- Task7 (W20, "Document Structure & Data Extraction Specialist"): dynamic header/
# metadata boundary detection — employee metadata (ชื่อ-สกุล/สังกัดแผนก) printed above the
# real data table must not be folded into the table body as if it were part of the header/
# data rows, and the header row isn't always row 1. ---

def test_load_manual_bytes_xlsx_splits_metadata_rows_above_detected_header():
    content = _xlsx_bytes(("Data", [
        ["ชื่อ-สกุล", "สมชาย ใจดี"],
        ["สังกัดแผนก", "IT"],
        ["วันที่", "รายละเอียด", "เวลาทำงาน", "จำนวนชั่วโมง"],
        ["2026-07-01", "ประชุมทีม", "09:00-17:00", 8],
    ]))
    text = load_manual_bytes(content, "timesheet.xlsx")
    lines = [line for line in text.split("\n") if line]
    assert lines == [
        "# Sheet: Data",
        "## Document Metadata",
        # sheet.iter_rows() reads the full rectangular used range (max_column across every
        # row = 4, from the header/data rows below) — these shorter 2-cell metadata rows
        # pad out to 4 cells with trailing empty strings, same as any other short row would
        # (see test_load_manual_bytes_xlsx_none_cells_become_empty_string above).
        "ชื่อ-สกุล | สมชาย ใจดี |  | ",
        "สังกัดแผนก | IT |  | ",
        "## Table",
        "วันที่ | รายละเอียด | เวลาทำงาน | จำนวนชั่วโมง",  # header row itself: plain, defines the labels
        "วันที่: 2026-07-01 | รายละเอียด: ประชุมทีม | เวลาทำงาน: 09:00-17:00 | จำนวนชั่วโมง: 8",
    ]


def test_load_manual_bytes_xlsx_no_header_keyword_leaves_sheet_unchanged():
    """ไม่มีแถวไหนตรงกับ header keyword ที่รู้จักเลย (เช่นตารางราคาสินค้าทั่วไป) — ต้องไม่ใส่
    "## Document Metadata"/"## Table" แทรกเข้ามาเลย และไม่ label คอลัมน์เพิ่ม (คงรูปแบบเดิม
    positional "a | b" ทุกประการ — feature นี้ scope เฉพาะ sheet ที่ตรงกับ known header keywords)"""
    content = _xlsx_bytes(("Data", [["Name", "Price"], ["Widget", 9.99]]))
    text = load_manual_bytes(content, "report.xlsx")
    assert "## Document Metadata" not in text
    assert "## Table" not in text
    assert "Name: " not in text
    lines = [line for line in text.split("\n") if line]
    assert lines == ["# Sheet: Data", "Name | Price", "Widget | 9.99"]


def test_load_manual_bytes_xlsx_header_on_first_row_still_labels_data_rows():
    """header keyword เจอที่แถวแรกพอดี (ไม่มี metadata แทรกอยู่ก่อนหน้าให้ต้องแยก section) —
    ไม่ต้องมี "## Document Metadata"/"## Table" (ไม่มีอะไรให้แยก) แต่แถวข้อมูลยังต้องถูก label
    ชื่อคอลัมน์กำกับอยู่ดี (ดู test_load_manual_bytes_xlsx_daily_detail_column_stays_labeled_
    and_readable_across_many_rows ด้านล่างสำหรับเหตุผลเต็ม — บั๊กจริงที่ user รายงาน)"""
    content = _xlsx_bytes(("Data", [
        ["วันที่", "รายละเอียด", "เวลาทำงาน", "จำนวนชั่วโมง"],
        ["2026-07-01", "ประชุมทีม", "09:00-17:00", 8],
    ]))
    text = load_manual_bytes(content, "timesheet.xlsx")
    assert "## Document Metadata" not in text
    assert "## Table" not in text
    lines = [line for line in text.split("\n") if line]
    assert lines == [
        "# Sheet: Data",
        "วันที่ | รายละเอียด | เวลาทำงาน | จำนวนชั่วโมง",
        "วันที่: 2026-07-01 | รายละเอียด: ประชุมทีม | เวลาทำงาน: 09:00-17:00 | จำนวนชั่วโมง: 8",
    ]


def test_load_manual_bytes_xlsx_daily_detail_column_stays_labeled_and_readable_across_many_rows():
    """บั๊กจริงที่ user รายงาน (พร้อมสกรีนช็อตไฟล์จริง): LLM ยืนยันหนักแน่นว่าคอลัมน์
    "รายละเอียดการฝึกงาน" "ไม่มีข้อมูล"/"ทุกแถวเว้นว่างไว้" ทั้งที่ในไฟล์จริงมีข้อความ (Setup,
    Oic claim, ...) อยู่ครบทุกแถว — สาเหตุที่น่าจะเป็นไปได้คือ bare positional "a | b | c" text
    ที่ยาวหลายสิบแถวทำให้โมเดลนับตำแหน่งคอลัมน์ผิดพลาด ต้อง label ชื่อคอลัมน์กำกับไว้ที่ทุกค่า
    โดยตรง (ไม่ต้องนับตำแหน่ง) — จำลองโครงสร้างเดียวกับไฟล์จริง (header 1 แถว, ตามด้วยหลายแถว
    ข้อมูลที่มีคอลัมน์ "รายละเอียดการฝึกงาน" ไม่ว่างเปล่าเลยสักแถว)"""
    content = _xlsx_bytes(("Data", [
        ["วันที่", "รายละเอียดการฝึกงาน", "เวลาทำงาน", "ชั่วโมง"],
        ["1 ก.ค. 69", "Setup", "9.00-18.00", 8],
        ["2 ก.ค. 69", "Oic claim", "9.00-18.00", 8],
        ["3 ก.ค. 69", "Oic claim", "9.00-18.00", 8],
        ["6 ก.ค. 69", "Oic claim", "9.00-18.00", 8],
    ]))
    text = load_manual_bytes(content, "timesheet.xlsx")
    lines = [line for line in text.split("\n") if line]
    # ทุกแถวข้อมูลต้องมี "รายละเอียดการฝึกงาน: <ค่าจริง>" ปรากฏชัดเจน ไม่ใช่แค่ค่าดิบลอยๆ ที่ต้อง
    # นับตำแหน่งเทียบ header เอาเอง
    assert "รายละเอียดการฝึกงาน: Setup" in lines[2]
    assert "รายละเอียดการฝึกงาน: Oic claim" in lines[3]
    assert "รายละเอียดการฝึกงาน: Oic claim" in lines[4]
    assert "รายละเอียดการฝึกงาน: Oic claim" in lines[5]
    # และต้องกำกับคู่กับ "วันที่" ของแถวเดียวกันด้วย (ไม่ใช่แค่ label เดี่ยวๆ ลอยไม่มีบริบทวันที่)
    assert lines[2].startswith("วันที่: 1 ก.ค. 69 |")
    assert lines[3].startswith("วันที่: 2 ก.ค. 69 |")


def test_load_manual_bytes_xlsx_merged_title_row_does_not_get_picked_as_header():
    """บั๊กจริงที่ user รายงาน (ไล่ debug ด้วยไฟล์จริงของ user ตรงๆ): แถวหัวเอกสาร (title) ที่
    merge ครอบทั้งแถว เช่น "ใบลงเวลาทำงานนักศึกษาฝึกงาน" บังเอิญมีคำว่า "เวลาทำงาน" ปนอยู่เป็น
    substring — ถ้าเจอ keyword แค่ตัวเดียวก็ถือว่าเป็น header แล้ว แถว title นี้จะถูกเข้าใจผิดว่า
    เป็น header แทนที่จะเป็น header จริงที่อยู่ถัดไปอีก 2 แถว ทำให้ label ของทุกคอลัมน์ในตารางทั้ง
    หมดผิดเพี้ยนไปหมด (ทุกค่ากลายเป็น "ใบลงเวลาทำงานนักศึกษาฝึกงาน: ..." แทนชื่อคอลัมน์จริง เช่น
    "รายละเอียดการฝึกงาน") — ต้องเจอ keyword อย่างน้อย 2 ตัวที่ "ต่างกัน" ถึงจะนับเป็น header แถว
    title ที่ merge ค่าเดียวซ้ำทั้งแถวจะแมตช์ได้แค่ keyword เดียวเสมอ (ไม่นับเป็น header)"""
    content = _xlsx_bytes(("Data", [
        ["ใบลงเวลาทำงานนักศึกษาฝึกงาน"] * 6,  # title row — merged in the real file, "ใบลงเวลา"
                                                # part of "เวลาทำงาน" false-matches 1 keyword
        ["ชื่อ-สกุล", "สยามยุทธ์ ผาสีดา"],
        ["วันที่", "รายละเอียดการฝึกงาน", "เวลาทำงาน", "จาก", "ถึง", "ชั่วโมง"],
        ["2026-07-01", "Setup", "", "9.00 น.", "18.00 น.", 8],
    ]))
    text = load_manual_bytes(content, "timesheet.xlsx")
    lines = [line for line in text.split("\n") if line]
    assert lines == [
        "# Sheet: Data",
        "## Document Metadata",
        "ใบลงเวลาทำงานนักศึกษาฝึกงาน | ใบลงเวลาทำงานนักศึกษาฝึกงาน | ใบลงเวลาทำงานนักศึกษาฝึกงาน | "
        "ใบลงเวลาทำงานนักศึกษาฝึกงาน | ใบลงเวลาทำงานนักศึกษาฝึกงาน | ใบลงเวลาทำงานนักศึกษาฝึกงาน",
        "ชื่อ-สกุล | สยามยุทธ์ ผาสีดา |  |  |  | ",
        "## Table",
        "วันที่ | รายละเอียดการฝึกงาน | เวลาทำงาน | จาก | ถึง | ชั่วโมง",
        "วันที่: 2026-07-01 | รายละเอียดการฝึกงาน: Setup | เวลาทำงาน:  | จาก: 9.00 น. | ถึง: 18.00 น. | ชั่วโมง: 8",
    ]


# --- Excel extractor: .csv support (spec ระบุ .xlsx / .csv ทั้งคู่ แต่ก่อนหน้านี้
# load_manual_bytes() รองรับแค่ .xlsx เท่านั้น) ---

def test_load_manual_bytes_csv_parses_rows_pipe_delimited():
    content = "Name,Date,Hours\nSomchai,2026-07-01,8\nSomsri,2026-07-02,6\n".encode("utf-8")
    text = load_manual_bytes(content, "timesheet.csv")
    lines = [line for line in text.split("\n") if line]
    assert lines == ["Name | Date | Hours", "Somchai | 2026-07-01 | 8", "Somsri | 2026-07-02 | 6"]


def test_load_manual_bytes_csv_skips_fully_empty_rows():
    content = "A,B\n1,2\n,\n3,4\n".encode("utf-8")
    text = load_manual_bytes(content, "data.csv")
    lines = [line for line in text.split("\n") if line]
    assert lines == ["A | B", "1 | 2", "3 | 4"]


def test_load_manual_bytes_csv_handles_quoted_commas_in_fields():
    content = 'Name,Note\nSomchai,"Oic claim, urgent"\n'.encode("utf-8")
    text = load_manual_bytes(content, "data.csv")
    assert 'Somchai | Oic claim, urgent' in text


def test_load_manual_bytes_csv_strips_bom_if_present():
    content = "Name,Hours\nSomchai,8\n".encode("utf-8-sig")
    text = load_manual_bytes(content, "data.csv")
    lines = [line for line in text.split("\n") if line]
    assert lines[0] == "Name | Hours"
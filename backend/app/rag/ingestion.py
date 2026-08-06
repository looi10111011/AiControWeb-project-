"""Manual ingestion: PDF/DOCX/TXT -> chunk -> embed -> เก็บใน ChromaDB.
W1: skeleton only. W3: implement จริง.
W3: เปลี่ยนมาใช้ local embedding (all-MiniLM-L6-v2 ผ่าน ChromaDB) แทน Gemini API
    — collection.upsert() ส่ง documents ดิบไป ให้ ChromaDB embed ให้เองอัตโนมัติ
"""

from io import BytesIO
from pathlib import Path
import csv
import hashlib
import re
from typing import List

# สำหรับอ่านไฟล์
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

from backend.app.rag.chroma_client import get_collection

# จบประโยคด้วย . ! ? (ตามด้วยเว้นวรรค) หรือขึ้นบรรทัดใหม่ — ใช้หาจุดตัดที่ไม่ทำให้ประโยคขาด
_SENTENCE_BOUNDARY_RE = re.compile(r'(?<=[.!?])\s+|\n+')


def load_manual(path: Path) -> str:
    """โหลดไฟล์ PDF/DOCX/TXT เป็น text"""
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")
    
    ext = path.suffix.lower()
    
    if ext == '.txt':
        with open(path, 'r', encoding='utf-8') as f:
            return f.read()
    
    elif ext == '.docx':
        doc = Document(path)
        text = '\n'.join([para.text for para in doc.paragraphs])
        return text
    
    elif ext == '.pdf':
        reader = PdfReader(path)
        text = ''
        for page in reader.pages:
            text += page.extract_text() + '\n'
        return text

    else:
        raise ValueError(f"Unsupported file type: {ext}")



# Task7 (W20, "Document Structure & Data Extraction Specialist" — MODULE: DYNAMIC HEADER &
# DATA BOUNDARY DETECTION): the real column-header row of a Thai timesheet-style export
# isn't always row 1 — document metadata printed above the table (ชื่อ-สกุล/สังกัดแผนก, see
# _split_metadata_and_table() below) pushes it down. Scan for these known header labels to
# *find* the header row instead of assuming index 0.
_XLSX_HEADER_KEYWORDS = ("วันที่", "รายละเอียด", "เวลาทำงาน", "จำนวนชั่วโมง")


def _label_row_cells(header_values: List[str], row_values: List[str]) -> str:
    """Task8-follow-up (real bug user hit: the LLM confidently claimed a column that clearly
    had data on every row — see screenshot in W20 session — was "empty on every row"): a bare
    positional "value | value | value" line forces the reader to *count* pipe segments against
    the header row to know which value belongs to which column, which gets error-prone fast
    once a sheet has 20+ data rows and a header that itself spans two merged rows (like the
    "เวลาทำงาน" header here, merged over "จาก"/"ถึง" sub-columns) — pair every cell with its own
    column label explicitly instead, so a fact like "รายละเอียดการฝึกงาน: Setup" is unambiguous
    standing completely on its own, no column-counting required at all."""
    pairs = []
    for i, value in enumerate(row_values):
        label = header_values[i] if i < len(header_values) and header_values[i] else f"col{i + 1}"
        pairs.append(f"{label}: {value}")
    return ' | '.join(pairs)


def _split_metadata_and_table(rows_values: List[List[str]]) -> List[str]:
    """Task7 + Task8 follow-up: rows *above* the detected header row (first row matching
    _XLSX_HEADER_KEYWORDS) are document metadata (เช่น "ชื่อ-สกุล: สมชาย ใจดี", "สังกัดแผนก:
    IT") — split into their own "## Document Metadata" section instead of being folded into
    the table body as if they were extra header/data rows. Every row from the header onward
    (the header row itself stays as plain "a | b | c" — it's what *defines* the labels, not a
    fact that needs one) is rendered as explicit "column_label: value" pairs via
    _label_row_cells() instead of bare positional text, so column identity survives even in a
    wide/irregular table (see _label_row_cells() docstring for the bug this fixes).

    No row matches at least 2 *distinct* header keywords (most sheets — plain data tables with
    an ordinary header already on row 1, like a price list) -> plain positional "a | b | c"
    dump for every row, completely unchanged from before this rule existed.

    Real bug this guards against (found debugging a real user-reported file): a document
    title merged across the whole row (e.g. "ใบลงเวลาทำงานนักศึกษาฝึกงาน", forward-filled
    identically into every cell of that row by the merge) happens to contain "เวลาทำงาน" as a
    plain substring — with a "match >= 1 keyword" rule that title row won on being scanned
    *before* the real header 2 rows down, silently mislabeling every column for the entire
    rest of the table. A genuine header row always has multiple *different* keywords landing
    in different cells (วันที่ in one cell, รายละเอียด in another, ...); a single merged title
    only ever matches the same one keyword repeated verbatim across every cell — requiring 2+
    distinct keyword hits tells those two cases apart without needing to special-case merges
    or titles explicitly."""
    header_idx = next(
        (
            i for i, row in enumerate(rows_values)
            if len({kw for cell in row for kw in _XLSX_HEADER_KEYWORDS if kw in cell}) >= 2
        ),
        None,
    )
    if header_idx is None:
        return [' | '.join(row) for row in rows_values]

    header_values = rows_values[header_idx]
    metadata_lines = [' | '.join(row) for row in rows_values[:header_idx]]
    table_lines = [
        ' | '.join(header_values),
        *(_label_row_cells(header_values, row) for row in rows_values[header_idx + 1:]),
    ]
    if not metadata_lines:
        return table_lines
    return ["## Document Metadata", *metadata_lines, "## Table", *table_lines]


def _xlsx_bytes_to_text(content: bytes) -> str:
    """แปลง XLSX เป็น text แบบตาราง อ่านง่ายสำหรับ LLM (ไม่ใช่ data structure) — เรียงตาม
    sheet, ข้ามแถวที่ว่างทั้งแถว — คอลัมน์ของแถวข้อมูล (หลังเจอ header ที่รู้จัก ดู
    _split_metadata_and_table()) ถูก label ชื่อคอลัมน์กำกับไว้ให้ชัดเจนต่อ cell เลย ไม่ใช่แค่
    " | " คั่นตำแหน่งเฉยๆ (ดู _label_row_cells())
    data_only=True: อ่านค่าที่ cache ไว้ล่าสุดของ formula cell (ไม่ใช่สูตรดิบ) — ตรงกับที่
    user เห็นตอนเปิดไฟล์จริงใน Excel

    Excel extractor: merged cell (เช่น ชื่อพนักงาน/แผนกที่ merge ครอบหลายแถววันที่ในไฟล์
    ลงเวลา) openpyxl เก็บค่าจริงไว้แค่ cell มุมบนซ้ายของแต่ละ merge range เท่านั้น — cell
    อื่นๆ ในช่วงเดียวกันเป็น MergedCell ที่ .value เป็น None เสมอแม้จะมองเห็นค่าใน Excel จริง
    ก็ตาม ต้อง forward-fill ค่าจาก top-left cell ให้ทุก cell ในช่วง merge เดียวกันก่อน ไม่งั้น
    sub-row ที่พึ่ง merge จาก parent จะกลายเป็นค่าว่างเปล่าไปเฉยๆ ทั้งที่ข้อมูลมีอยู่จริง"""
    workbook = load_workbook(BytesIO(content), data_only=True)
    sections = []
    for sheet in workbook.worksheets:
        merged_values = {}
        for merged_range in sheet.merged_cells.ranges:
            top_left_value = sheet.cell(merged_range.min_row, merged_range.min_col).value
            for r in range(merged_range.min_row, merged_range.max_row + 1):
                for c in range(merged_range.min_col, merged_range.max_col + 1):
                    merged_values[(r, c)] = top_left_value

        rows_values: List[List[str]] = []
        for row in sheet.iter_rows():
            values = [
                merged_values.get((cell.row, cell.column)) if cell.value is None else cell.value
                for cell in row
            ]
            if all(v is None for v in values):
                continue
            rows_values.append(['' if v is None else str(v) for v in values])
        if rows_values:
            lines = _split_metadata_and_table(rows_values)
            sections.append(f"# Sheet: {sheet.title}\n" + '\n'.join(lines))
    return '\n\n'.join(sections)


def _csv_bytes_to_text(content: bytes) -> str:
    """แปลง CSV เป็น text แบบตาราง รูปแบบเดียวกับ _xlsx_bytes_to_text() ด้านบน (cell คั่นด้วย
    " | ", ข้ามแถวที่ว่างทั้งแถว) ให้ LLM อ่านสม่ำเสมอไม่ว่าไฟล์จะเป็น .xlsx หรือ .csv — ใช้
    csv.reader() (ไม่ใช่ split(",") มือ) เพื่อรองรับ field ที่มี comma/quote อยู่ในค่าเองถูก
    ต้องตามสเปค CSV จริง utf-8-sig: ตัด BOM ทิ้งถ้ามี (Excel ใส่ BOM มาด้วยเวลา export CSV)"""
    text = content.decode('utf-8-sig')
    rows_text = []
    for row in csv.reader(text.splitlines()):
        if not row or all(not cell.strip() for cell in row):
            continue
        rows_text.append(' | '.join(row))
    return '\n'.join(rows_text)


def load_manual_bytes(content: bytes, filename: str) -> str:
    """เหมือน load_manual() ด้านบนแต่รับ bytes ตรงๆ ไม่ต้องเขียนลงดิสก์ก่อน — ใช้กับไฟล์ที่
    user แนบเข้ามาผ่าน API โดยตรง (ต่างจาก load_manual ที่ใช้กับ manual ที่มีอยู่บนดิสก์แล้ว
    ของ RAG ingestion pipeline เดิม) ไม่มี temp file ให้ต้อง cleanup และไม่มีช่องโหว่ path
    traversal จาก filename ที่ user ตั้งเอง (ใช้แค่ดู extension ไม่เคยใช้เป็น path จริง)"""
    ext = Path(filename).suffix.lower()

    if ext == '.txt':
        return content.decode('utf-8')

    elif ext == '.docx':
        doc = Document(BytesIO(content))
        return '\n'.join([para.text for para in doc.paragraphs])

    elif ext == '.pdf':
        reader = PdfReader(BytesIO(content))
        text = ''
        for page in reader.pages:
            text += page.extract_text() + '\n'
        return text

    elif ext == '.xlsx':
        return _xlsx_bytes_to_text(content)

    elif ext == '.csv':
        return _csv_bytes_to_text(content)

    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _tail_at_word_boundary(text: str, size: int) -> str:
    """เอาท้ายข้อความมาไม่เกิน size ตัวอักษร แต่ขยับไปขอบเขตคำ (เว้นวรรค) แรกที่เจอ
    กันคำถูกตัดครึ่งตอนใช้เป็น overlap ต่อท้าย chunk ถัดไป"""
    if len(text) <= size:
        return text
    tail = text[-size:]
    space_idx = tail.find(' ')
    return tail[space_idx + 1:] if space_idx != -1 else tail


def _split_by_words(text: str, size: int) -> List[str]:
    """ตัดข้อความยาวตามขอบเขตคำ (เว้นวรรค) แทนการตัดกลางคำตรงๆ ตามจำนวนตัวอักษร
    ใช้เป็น fallback สุดท้ายเมื่อประโยคเดียวก็ยังยาวเกิน chunk_size (พบยาก)"""
    words = text.split(' ')
    pieces = []
    current = ''
    for word in words:
        candidate = f'{current} {word}'.strip() if current else word
        if len(candidate) > size and current:
            pieces.append(current)
            current = word
        else:
            current = candidate
    if current:
        pieces.append(current)

    # กรณีคำเดียวยาวเกิน size จริงๆ (แทบไม่เกิดขึ้น) ค่อย hard-split ตามตัวอักษร
    result = []
    for piece in pieces:
        if len(piece) <= size:
            result.append(piece)
        else:
            result.extend(piece[i:i + size] for i in range(0, len(piece), size))
    return result


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    """แบ่งข้อความเป็น chunks ตามย่อหน้า (บรรทัดว่างคั่น) — แต่ละหัวข้อ/ย่อหน้า
    จะไม่ถูกรวมเข้ากับหัวข้ออื่นในบรรทัดเดียวกัน กัน chunk ใหญ่เกินไปแบบไม่จำเป็น
    ถ้าย่อหน้าไหนยาวเกิน chunk_size เอง จะแบ่งตามขอบเขตประโยค (ไม่ตัดกลางประโยค/กลางคำ)
    """
    if not text:
        return []

    paragraphs = [p.strip() for p in re.split(r'\n\s*\n', text) if p.strip()]
    if not paragraphs:
        return []

    chunks: List[str] = []
    current = ''

    for para in paragraphs:
        # ย่อหน้าพอดีขนาด -> เป็น chunk แยกของตัวเอง ไม่รวมกับย่อหน้าอื่น
        units = [para] if len(para) <= chunk_size else _SENTENCE_BOUNDARY_RE.split(para)

        for unit in units:
            unit = unit.strip()
            if not unit:
                continue

            # ประโยคเดียวก็ยังยาวเกิน chunk_size (พบยาก) -> ตัดตามขอบเขตคำแทน
            pieces = [unit] if len(unit) <= chunk_size else _split_by_words(unit, chunk_size)

            for piece in pieces:
                if current and len(current) + len(piece) + 1 > chunk_size:
                    chunks.append(current)
                    # overlap: เอาท้ายประโยคสุดท้ายของ chunk ก่อนหน้ามาต่อ ไม่ตัดกลางคำ
                    current = _tail_at_word_boundary(current, overlap).strip() if overlap > 0 else ''
                    current = f'{current} {piece}'.strip() if current else piece
                else:
                    current = f'{current} {piece}'.strip() if current else piece

        # จบย่อหน้าแล้ว ตัด chunk ทันที ไม่ดึงย่อหน้าถัดไปมารวม
        if current:
            chunks.append(current)
            current = ''

    if current:
        chunks.append(current)

    return chunks


def ingest_manual(path: Path):
    """load -> chunk -> add to collection (ChromaDB embed ให้เองด้วย local model)"""
    print(f"📄 Loading: {path}")

    # 1. Load text
    text = load_manual(path)
    print(f"   ✅ Loaded {len(text)} characters")

    # 2. Chunk
    chunks = chunk_text(text)
    print(f"   ✅ Created {len(chunks)} chunks")

    if not chunks:
        print("⚠️ No chunks created, skipping...")
        return

    # 3. Add to ChromaDB — ไม่ embed เองแล้ว ส่ง documents ดิบไปให้ collection
    #    embed ด้วย local model (all-MiniLM-L6-v2) ที่ผูกไว้กับ collection ใน chroma_client.py
    #    (โหลดโมเดลครั้งแรก ~90MB แล้ว cache ไว้ใช้ครั้งต่อไป)
    print("   🔧 Generating embeddings (local all-MiniLM-L6-v2)...")
    collection = get_collection()

    # สร้าง ids, documents, metadatas
    ids = []
    documents = []
    metadatas = []

    for i, chunk in enumerate(chunks):
        # สร้าง id จาก hash ของชื่อไฟล์ + index (deterministic -> ingest ไฟล์เดิมซ้ำแล้ว upsert ทับของเดิมได้)
        doc_id = hashlib.md5(f"{path.stem}_{i}".encode()).hexdigest()
        ids.append(doc_id)
        documents.append(chunk)
        metadatas.append({
            "source": path.name,
            "chunk_index": i,
            "total_chunks": len(chunks),
            "path": str(path)
        })

    # upsert แทน add: กัน error/ถูกข้ามเงียบๆ ตอน ingest ไฟล์เดิมซ้ำ (เช่น อัปเดตคู่มือ)
    collection.upsert(
        documents=documents,
        metadatas=metadatas,
        ids=ids
    )

    print(f"✅ Done! Added {len(chunks)} chunks to collection '{collection.name}'")